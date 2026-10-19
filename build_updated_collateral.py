"""Build the executive deck and email from reviewed, structured content.

Run from this repository root: python3 build_updated_collateral.py
Uses native PowerPoint bullet paragraphs. Keeps the first pre-update copies in archive.
"""
from pathlib import Path
from html import escape
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.shapes import MSO_SHAPE
from PIL import ImageFont
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parent
NAVY, ORANGE, INK, MUTED, PANEL = '0A2540', 'F38020', '334155', '64748B', 'F3F6FA'
FONT = 'DejaVu Sans'
FONT_FILE = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
BOLD_FILE = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
DATE = '19 October 2026'
# Each slide has two clear bullet groups; these are editable PowerPoint text boxes.
SLIDES = [
('VulnSuite', 'OPERATIONS RELEASE 0.4',
 ('Assessment and detection', ['Code, dependencies and infrastructure assessment', 'Linux and Windows security telemetry', 'Malware signals and cross-host correlation']),
 ('Security operations', ['Verified identity, roles and a shared SOC view', 'Incident ownership and central agent enrollment', 'Reviewed quarantine, isolation and cloud-WAF response']),
 'Implemented source with automated validation | Production rollout requires deployment-specific staging.'),
('What the suite addresses', 'SECURITY OUTCOMES',
 ('Reduce exposure', ['Find vulnerable software and unsafe configurations', 'Assess exposed services, APIs and TLS settings', 'Prioritise findings using severity and asset context']),
 ('Investigate and respond', ['Connect host alerts with vulnerability findings', 'Assign incidents, review evidence and track action', 'Approve bounded containment and inspect receipts']),
 'Selected detection signals support investigation; this release does not claim complete EDR coverage.'),
('Platform coverage', 'ASSESSMENT + OPERATIONS',
 ('Assessment modules', ['SCA, secrets and static code analysis', 'Network, cloud, containers and asset discovery', 'IaC, Kubernetes and external attack surface', 'API testing and software supply-chain checks']),
 ('Protection and operations', ['Linux audit and Windows event adapters', 'Access, process, malware and integrity signals', 'SOC console, investigations and agent lifecycle', 'Signed notifications and approved response plans']),
 'Assessment and protection services retain separate persistence and deployment boundaries.'),
('Software and supply-chain risk', 'APPLICATION SECURITY',
 ('Code and dependencies', ['Trivy, pip-audit and OSV-Scanner collectors', 'Semgrep and Bandit static analysis', 'Gitleaks and TruffleHog secret detection']),
 ('Build and artifact evidence', ['Syft software bill of materials', 'Cosign signature and attestation checks', 'VEX statements captured in finding evidence', 'Risk scoring and remediation lifecycle tracking']),
 'Coverage depends on configured targets, installed tools and available advisory data.'),
('Infrastructure and API assessment', 'EXPOSURE MANAGEMENT',
 ('Infrastructure', ['Nmap service discovery and testssl.sh TLS checks', 'Prowler cloud configuration collectors', 'Checkov / Trivy IaC and Kubernetes posture checks']),
 ('Applications and APIs', ['Nuclei and OWASP ZAP assessment collectors', 'Target allowlists and active-scan configuration gates', 'External discovery with Subfinder and httpx']),
 'Assessment scans complement host monitoring and inline WAF enforcement.'),
('Verified identity and access', 'SSO TOKEN VERIFICATION + RBAC',
 ('Authentication and roles', ['Verify token signature, issuer, audience and expiry', 'Require signed tenant and explicit role claims', 'Separate viewer, auditor, analyst and admin rights', 'Use external identity-provider sign-in and MFA']),
 ('Tenant and approval boundaries', ['Bind assessment transactions to signed tenant ID', 'Enforce PostgreSQL row-level security policies', 'Bind protection access to one organisation', 'Require a different admin to approve playbooks']),
 'JWT verification and RLS controls are implemented; live IdP integration and database-role enforcement need staging.'),
('Unified operations console', 'ALERTS + FINDINGS + AGENTS',
 ('Shared investigation view', ['View intrusion alerts and vulnerability findings', 'Inspect agent health, sensor status and version drift', 'Create and assign incidents from selected alerts', 'Review response plans before approval']),
 ('Operational safeguards', ['Apply the same API role checks to console actions', 'Keep access tokens in browser-tab memory', 'Report unavailable upstream sources explicitly', 'Retrieve findings through a fixed HTTPS connector']),
 'The console shows bounded recent lists; it is not an exhaustive cross-platform evidence search.'),
('Broader endpoint telemetry', 'WINDOWS EVENTS + LINUX AUDIT',
 ('Windows monitoring', ['Security log authentication and network logons', 'RDP and SMB evidence where source events support it', 'Sysmon process-creation records', 'Selected fields without raw command-line capture']),
 ('Linux monitoring', ['Event-driven audit execution records', 'Existing /proc process polling as a complement', 'SSH access, API denials and file-integrity signals', 'Sensor health and missing-telemetry indicators']),
 'Windows now includes monitoring and approved quarantine, restoration, host isolation and release. Native Windows staging validation is still required.'),
('Process and malware detection', 'SELECTED HOST SIGNALS',
 ('Suspicious activity', ['Recognised web-server parent spawning a shell', 'Execution from writable temporary locations', 'Deleted executable signals on Linux', 'Event-driven records improve short-process visibility']),
 ('Malware evidence', ['Known-bad SHA256 matches on selected targets', 'Optional ClamAV scans of selected local files', 'Path, signature and hash retained for review', 'Unavailable scanners surface degraded coverage']),
 'No complete memory scanning, full-disk protection or automatic process-termination claim is made.'),
('Lateral-movement signals', 'CROSS-HOST CORRELATION',
 ('Correlated patterns', ['Password spraying across multiple hosts', 'Same-user successful-login fan-out', 'Suspected pivots from already flagged hosts', 'Remote-access evidence from supported event sources']),
 ('Investigation context', ['Administrator-defined host-to-IP ownership', 'Time-bounded correlation with retained evidence', 'Retry, out-of-order and restart handling', 'Expiring exceptions for reviewed false positives']),
 'NAT, shared accounts and legitimate automation need review; protocol labels require supporting source evidence.'),
('WAF integration', 'MODSECURITY + CLOUDFLARE',
 ('On-premises or proxy WAF', ['Nginx / ModSecurity and OWASP CRS templates', 'Sanitised serial-JSON rule-match collection', 'DetectionOnly starting configuration', 'No request-body or credential export']),
 ('Cloudflare response connector', ['Use a configured zone and custom-ruleset scope', 'Apply an approved source-IP block', 'Embed fixed expiry in the provider rule', 'Retry cleanup without extending the deadline']),
 'Rule ordering, provider permissions, propagation and real traffic enforcement remain staging checks.'),
('Reviewed local containment', 'APPROVAL + BOUNDED ENFORCEMENT',
 ('Quarantine and restoration', ['Use approved malware path and SHA256 evidence', 'Linux: private same-filesystem quarantine vault', 'Windows: file-handle checks and protected ACLs', 'Restore without overwriting existing files']),
 ('Isolation and release', ['Require separate approval and local policy', 'Linux: kernel-expiring nftables isolation', 'Windows: firewall rules with local recovery tasks', 'Preserve configured management access']),
 'Windows expiry uses local recovery tasks; Linux uses kernel timeouts. Native enforcement and recovery require staging.'),
('How the components fit', 'CURRENT ARCHITECTURE',
 ('Assessment service', ['Scanner collectors and orchestration workers', 'Canonical findings and deterministic risk scoring', 'Verified API access and PostgreSQL tenant policies', 'Read-only findings feed into the operations console']),
 ('Protection service', ['Authenticated Linux and Windows agents', 'One organisation per SQLite operational store', 'Detection, incidents, policies and response history', 'Independent delivery and cloud-response worker']),
 'SQLite protection persistence is not a horizontally scalable or highly available PostgreSQL backend.'),
('Central agent lifecycle', 'ENROLLMENT + POLICY + HEALTH',
 ('Credential management', ['Issue expiring, one-time enrollment credentials', 'Store hashes of enrolled host credentials', 'Rotate keys and revoke agent access centrally', 'Cancel unexecuted responses after revocation']),
 ('Inventory and policy', ['Track platform, agent version and queue health', 'Record administrator policy-version history', 'Flag drift from the desired software version', 'Keep privileged response enablement local to hosts']),
 'Agent software upgrades use an administrator-controlled rollout; this release does not remotely execute upgrades.'),
('Incident and notification workflows', 'OWNERSHIP + EVIDENCE + DELIVERY',
 ('Investigation workflow', ['Link alerts to owned incident cases', 'Track status, notes and evidence timelines', 'Protect concurrent edits with case versions', 'Review alerts and add scoped, expiring exceptions']),
 ('SIEM and ticketing bridges', ['Send signed JSON to fixed HTTPS endpoints', 'Persist delivery state, leases and retry attempts', 'Attach stable keys for receiver deduplication', 'Audit outcomes without storing upstream secrets']),
 'Webhooks support organisation-operated adapters; native Jira, ServiceNow or Splunk connectors are not implied.'),
('Reliable delivery and retention', 'RECOVERY + MAINTENANCE',
 ('Telemetry continuity', ['Keep events until exact acknowledgement', 'Replay retained events after service outages', 'Pause collection visibly when the queue is full', 'Archive older events without fresh containment']),
 ('Evidence maintenance', ['Create consistent SQLite backups and verify restore', 'Preview retention cleanup before applying it', 'Preserve incident and response-linked evidence', 'Create offline-agent and stale-sensor alerts']),
 'Schedule maintenance explicitly and measure capacity, recovery objectives and backup security in the target environment.'),
('Implementation validation', 'AUTOMATED SECURITY EVIDENCE',
 ('Security and protection tests', ['Signed-token rejection, roles and tenant controls', 'Enrollment, revocation and case concurrency', 'Windows/audit parsing and outage handling', 'Durable notification retries and retention behavior']),
 ('Workflow validation', ['Separate-actor approval through execution receipt', 'Linux/Windows quarantine and restore workflows', 'Mocked cloud-WAF expiry, cleanup and revocation', 'Console interaction and synthetic detector scenarios']),
 'Automated security/protection and end-to-end tests exercise code paths; mocked enforcement does not establish live efficacy.'),
('Required staging validation', 'DEPLOYMENT GATES',
 ('Identity and host environment', ['External identity-provider issuance and revocation', 'PostgreSQL RLS with the deployed runtime role', 'Windows Security/Sysmon and Linux audit collection', 'ClamAV installation, signatures and scan results']),
 ('Enforcement and workload', ['Cloudflare permissions, expiry and rule ordering', 'ModSecurity / CRS tuning and real intervention', 'Linux/Windows isolation, expiry and rollback', 'False positives, overhead, backlog and recovery']),
 'No production IdP changes, external notifications, cloud-WAF rules or live host firewall changes were made here.'),
('Further enhancement opportunities', 'NEXT ENGINEERING PRIORITIES',
 ('Scale and endpoint depth', ['Benchmark real organisation workloads', 'Design operational HA and enterprise persistence', 'Extend process, memory and network visibility', 'Validate Windows fleet coverage and recovery']),
 ('Integration and rollout maturity', ['Add provider-specific WAF and ticket adapters', 'Expand labelled detection evaluation datasets', 'Strengthen signed-package upgrade orchestration', 'Measure production coverage and response quality']),
 'These are future extensions; central enrollment, verified access, the SOC console and reviewed playbooks are now implemented.'),
('Proposed 30-day pilot', 'VALIDATE BEFORE PRODUCTION ROLLOUT',
 ('Weeks 1–2: establish coverage', ['Confirm owners, asset scope and identity integration', 'Enroll selected Linux and Windows staging hosts', 'Enable native audit sources and stage the WAF', 'Baseline normal activity and review noisy signals']),
 ('Weeks 3–4: validate and decide', ['Exercise incident and signed-notification workflows', 'Validate quarantine, isolation, expiry and rollback', 'Test outage replay, retention and verified recovery', 'Agree measurable production acceptance criteria']),
 'Pilot scope and dates are proposals; rollout follows evidence from staging and an approved operational runbook.'),
]

SUBJECT = 'VulnSuite 0.4 update — security operations, broader server monitoring and approved containment'
SECTIONS = [
('Implemented security operations', [
 'Verified identity and roles: access-token signature, issuer, audience, expiry and tenant checks; viewer/auditor, analyst and administrator permissions; assessment PostgreSQL row-level security controls.',
 'Unified SOC console: intrusion alerts, vulnerability findings, agent/sensor health, version drift and response review; findings use a fixed read-only HTTPS connection.',
 'Incident workflows: ownership, linked evidence, timelines, concurrent-edit protection, alert review and narrowly scoped exceptions with expiry.',
 'Central agent management: one-time enrollment, credential rotation/revocation, policy-version history and desired-version tracking; software upgrades remain administrator controlled.'
]),
('Expanded monitoring and response', [
 'Endpoint telemetry: Windows Security/Sysmon adapters and event-driven Linux audit records complement SSH/API abuse, file-integrity and process monitoring.',
 'Windows now includes monitoring and approved quarantine, restoration, host isolation and release. Native Windows staging validation is still required.',
 'Detection: suspicious process behavior, known-bad SHA256 matches, optional selected-file ClamAV scans and broader cross-host login/pivot correlation.',
 'WAF integration: ModSecurity/OWASP CRS configuration and sanitised telemetry, plus a configured Cloudflare connector with provider-evaluated block expiry and cleanup retries.',
 'Controlled Windows and Linux containment: separate-actor approval, local policy gates, hash-verified quarantine/restoration and temporary host isolation/release. Windows uses protected file ACLs and local firewall recovery tasks; Linux uses nftables timeouts.'
]),
('Reliability and operational integrations', [
 'Reliable delivery: persistent queues, exact acknowledgements, visible backpressure and historical-event retention; old evidence cannot initiate fresh containment.',
 'Evidence maintenance: consistent SQLite backup, verified restoration, previewable retention cleanup and preservation of incident/response-linked evidence.',
 'SIEM/ticketing bridges: signed HTTPS JSON webhooks with durable retries and delivery IDs; receiving adapters must verify signatures and deduplicate events.',
 'Operational health: offline-agent and stale-sensor alerts, queue health, audited actions and execution receipts.'
]),
('Validation status and remaining deployment work', [
 'Automated security/protection and end-to-end tests cover access boundaries, enrollment, detection, outage recovery, approval, quarantine/restoration and mocked provider/firewall response.',
 'Required staging: external identity-provider issuance/revocation, PostgreSQL RLS under the deployed runtime role, Windows/auditd collection, Windows file ACLs and recovery tasks, ClamAV signatures, WAF enforcement and native host isolation/rollback.',
 'The protection backend remains one organisation per SQLite store; horizontal scale, high availability and production capacity are not established.',
 'Coverage is selected monitoring and response, not a full EDR or memory-protection platform.',
 'The implementation did not send external notifications, change live WAF/firewall rules, deploy identity settings or upgrade host software.'
]),
('Proposed pilot and subsequent enhancements', [
 'Run a 30-day pilot on representative non-production Linux and Windows hosts, using the organisation’s identity provider and staging WAF.',
 'Exercise case management, signed notifications, containment, expiry, rollback, outage replay and verified recovery; measure false positives, collection gaps and resource overhead.',
 'Use the results to prioritise operational HA/persistence, deeper endpoint/network visibility, additional vendor adapters and signed-package upgrade orchestration.',
 'Agree accountable owners, retention/recovery requirements and measurable acceptance criteria before production rollout.'
]),
('Updated material', [
 'VulnSuite_Executive_Deck_FINAL.pptx and VulnSuite_Executive_Deck_Review.pdf — structured executive briefing for the operations release.',
 'Operations guide — vulnsuite/protection/OPERATIONS.md; staging and maintenance — deploy/protection/OPERATIONS.md.'
])]


def color(value): return RGBColor.from_string(value)

def box(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size=18, fill=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.text = value
    p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color(fill)
    p.space_after = Pt(0)
    return shape


def wrapped(value, size, width):
    font = ImageFont.truetype(FONT_FILE, int(size * 4))
    words = value.split(); lines=[]; line=''
    for word in words:
        candidate=(line+' '+word).strip()
        if line and font.getlength(candidate)/4 > width*72:
            lines.append(line); line=word
        else: line=candidate
    if line: lines.append(line)
    return lines


def bullets(slide, items, x, y, w, h):
    shape = text(slide, '', x, y, w, h)
    tf=shape.text_frame
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=item; p.level=0
        p.font.name=FONT; p.font.size=Pt(18); p.font.color.rgb=color(INK)
        p.line_spacing=Pt(24); p.space_after=Pt(12)
        props=p._p.get_or_add_pPr(); props.set('marL',str(int(Inches(.20)))); props.set('indent',str(-int(Inches(.18))))
        bullet=OxmlElement('a:buChar'); bullet.set('char','•'); props.append(bullet)
    needed=sum(len(wrapped(item,18,w-.23))*24+12 for item in items)-12
    if needed > h*72: raise ValueError(f'Bullet overflow: {needed}pt > {h*72}pt')


def generate():
    archive=ROOT/'collateral/archive'; archive.mkdir(parents=True, exist_ok=True)
    for name in ['VulnSuite_Executive_Deck_FINAL.pptx','VulnSuite_Executive_Deck.pptx','VulnSuite_Executive_Deck_Review.pdf','VulnSuite_Intro_Email.txt','VulnSuite_Intro_Email.html']:
        backup=archive/(Path(name).stem+'_before_operations_0_3'+Path(name).suffix)
        if (ROOT/name).exists() and not backup.exists(): shutil.copy2(ROOT/name,backup)
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    prs.core_properties.title='VulnSuite 0.4 — assessment, operations and approved response'
    prs.core_properties.subject='Implemented operations release, automated validation and live staging gates'
    prs.core_properties.author='VulnSuite'
    pdfmetrics.registerFont(TTFont('VS',FONT_FILE)); pdfmetrics.registerFont(TTFont('VS-Bold',BOLD_FILE))
    pdf=canvas.Canvas(str(ROOT/'VulnSuite_Executive_Deck_Review.pdf'),pagesize=(13.333*72,7.5*72))
    pdf.setTitle('VulnSuite 0.4 — assessment, operations and approved response')
    pdf.setAuthor('VulnSuite')
    def pdftext(value,x,y,size=18,bold=False,fill=INK,width=None):
        pdf.setFillColor('#'+fill); pdf.setFont('VS-Bold' if bold else 'VS',size)
        lines=wrapped(value,size,width) if width else [value]
        for j,line in enumerate(lines): pdf.drawString(x*72,(7.5-y)*72-size-j*(size+6),line)
    for n,(title,kicker,left,right,note) in enumerate(SLIDES,1):
        s=prs.slides.add_slide(prs.slide_layouts[6])
        box(s,0,0,13.333,.10,ORANGE)
        text(s,kicker,.55,.35,12,.3,11,ORANGE,True)
        text(s,title,.55,.87,12.2,.65,34,NAVY,True)
        pdf.setFillColor('#'+ORANGE);pdf.rect(0,7.4*72,13.333*72,.1*72,fill=1,stroke=0)
        pdftext(kicker,.55,.35,11,True,ORANGE);pdftext(title,.55,.87,34,True,NAVY)
        for x,(heading,items) in zip([.55,6.87],[left,right]):
            box(s,x,1.85,5.91,4.6,PANEL)
            box(s,x,1.85,.05,4.6,ORANGE)
            text(s,heading,x+.27,2.1,5.36,.7,21,NAVY,True)
            bullets(s,items,x+.27,3.02,5.3,3.25)
            pdf.setFillColor('#'+PANEL);pdf.rect(x*72,(7.5-6.45)*72,5.91*72,4.6*72,fill=1,stroke=0)
            pdf.setFillColor('#'+ORANGE);pdf.rect(x*72,(7.5-6.45)*72,.05*72,4.6*72,fill=1,stroke=0)
            pdftext(heading,x+.27,2.1,21,True,NAVY,width=5.36)
            cursor=3.02
            for item in items:
                pdftext('•',x+.27,cursor,18,False,ORANGE)
                pdftext(item,x+.5,cursor,18,width=5.07)
                cursor+=(len(wrapped(item,18,5.07))*24+12)/72
        text(s,note,.55,6.65,12.2,.42,11,MUTED)
        text(s,f'CONFIDENTIAL • INTERNAL REVIEW  |  {DATE}',.55,7.13,11,.2,9,MUTED)
        text(s,f'{n:02d} / {len(SLIDES):02d}',12.08,7.13,.75,.2,9,MUTED)
        pdftext(note,.55,6.65,11,fill=MUTED,width=12.2)
        pdftext(f'CONFIDENTIAL • INTERNAL REVIEW  |  {DATE}',.55,7.13,9,fill=MUTED)
        pdftext(f'{n:02d} / {len(SLIDES):02d}',12.08,7.13,9,fill=MUTED)
        s.notes_slide.notes_text_frame.text='Scope: '+note+'\nImplementation and limitations: vulnsuite/protection/OPERATIONS.md\nStaging and maintenance: deploy/protection/OPERATIONS.md\nAutomated validation does not establish live enforcement or full EDR coverage.'
        pdf.showPage()
    pdf.save()
    prs.save(ROOT/'VulnSuite_Executive_Deck_FINAL.pptx')
    shutil.copy2(ROOT/'VulnSuite_Executive_Deck_FINAL.pptx',ROOT/'VulnSuite_Executive_Deck.pptx')
    intro='Please find the updated VulnSuite 0.4 briefing. The previously proposed operations enhancements are now implemented in source, with automated security and workflow validation. The structured summary below distinguishes delivered capabilities from the remaining live deployment checks.'
    lines=['Subject: '+SUBJECT,'','Classification: CONFIDENTIAL — INTERNAL USE ONLY','','Dear [CISO Name],','',intro,'']
    for i,(heading,items) in enumerate(SECTIONS,1):
        lines += [f'{i}. {heading.upper()}']+['   • '+item for item in items]+['']
    lines += ['Regards,','[Your Name]','[Your Role]','[Phone] | [Email]','','CONFIDENTIAL — INTERNAL USE ONLY']
    (ROOT/'VulnSuite_Intro_Email.txt').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    sections=''.join('<h2>'+str(i)+'. '+escape(h)+'</h2><ul>'+''.join('<li>'+escape(v)+'</li>' for v in items)+'</ul>' for i,(h,items) in enumerate(SECTIONS,1))
    page='''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>VulnSuite briefing</title><style>body{margin:0;background:#f3f6fa;color:#334155;font:16px/1.6 Arial,sans-serif}main{max-width:820px;margin:24px auto;background:white;padding:36px;border-top:6px solid #f38020}h1{color:#0a2540;font-size:26px;line-height:1.3}h2{color:#0a2540;font-size:19px;margin-top:28px}li{margin:9px 0}small{color:#64748b}ul{padding-left:25px}</style></head><body><main>'''
    page+='<small>CONFIDENTIAL — INTERNAL USE ONLY</small><h1>'+escape(SUBJECT)+'</h1><p>Dear [CISO Name],</p><p>'+escape(intro)+'</p>'+sections+'<p>Regards,<br>[Your Name]<br>[Your Role]<br>[Phone] | [Email]</p></main></body></html>'
    (ROOT/'VulnSuite_Intro_Email.html').write_text(page,encoding='utf-8')
    print(f'Generated {len(SLIDES)} slides, native bullet lists, TXT/HTML email and companion review PDF.')

if __name__=='__main__': generate()
